import pandas as pd
import os
import copy
import openpyxl

def load_excel(file_path):
    return pd.read_excel(file_path)


def save_excel(df, file_path):
    df.to_excel(file_path, index=False)


def prepare_result_columns(df):
    if "Status" not in df.columns:
        df["Status"] = ""

    if "Message" not in df.columns:
        df["Message"] = ""

    return df


def log_result(df, idx, status, message):
    df.loc[idx, "Status"] = status
    df.loc[idx, "Message"] = message



def add_log(df, idx, message):
    print(message)
    df.at[idx, "LOG"] += f"{message}\n"

def save_progress(input_file,df):
    temp_file = input_file.replace(".xlsx", "_temp.xlsx")
    df.to_excel(temp_file, index=False, engine="openpyxl")
    os.replace(temp_file, input_file)


def _copy_cell(src, dst):
    dst.value = src.value
    dst._style = copy.copy(src._style)


def unmerge_and_fill(input_file, output_file=None,
                     group_cols=("Parent Material", "Alternative BOM"),
                     fill_in_group_cols=("Base Qty after",),
                     component_col="Component after"):
    """
    Unmerge all merged cells and fill every row with its value (keeping format).
    Then fill blank group columns (e.g. Parent Material) from the row above,
    and fill blank fill_in_group_cols within the same Parent Material + Alt BOM.
    Saves to output_file (default: <input>_unmerged.xlsx). Input is not changed.
    """
    if output_file is None:
        output_file = input_file.replace(".xlsx", "_unmerged.xlsx")

    wb = openpyxl.load_workbook(input_file)
    ws = wb.active

    # 1. Unmerge and copy top-left value to every cell of the range
    ranges = list(ws.merged_cells.ranges)
    for r in ranges:
        top = ws.cell(r.min_row, r.min_col)
        value, style = top.value, top._style
        ws.unmerge_cells(str(r))
        for row in range(r.min_row, r.max_row + 1):
            for col in range(r.min_col, r.max_col + 1):
                cell = ws.cell(row, col)
                cell.value = value
                cell._style = copy.copy(style)
    print(f"Unmerged {len(ranges)} ranges")

    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    for name in (*group_cols, *fill_in_group_cols, component_col):
        if name not in headers:
            raise ValueError(f"Column '{name}' not found in {input_file}")
    comp_c = headers[component_col]

    # 2. Fill blank group columns from the row above (only rows that have a component)
    for name in group_cols:
        c = headers[name]
        for row in range(3, ws.max_row + 1):
            cell = ws.cell(row, c)
            if cell.value in (None, "") and ws.cell(row, comp_c).value not in (None, ""):
                _copy_cell(ws.cell(row - 1, c), cell)
                print(f"Row {row}: filled '{name}' = {cell.value} from row above")

    # 3. Fill blank columns within the same group
    group_c = [headers[n] for n in group_cols]
    for name in fill_in_group_cols:
        c = headers[name]
        key, first = None, None
        for row in range(2, ws.max_row + 1):
            row_key = tuple(ws.cell(row, g).value for g in group_c)
            cell = ws.cell(row, c)
            if row_key != key:
                key, first = row_key, cell
            elif cell.value in (None, ""):
                _copy_cell(first, cell)

    # 4. Report rows still missing required values
    missing = [row for row in range(2, ws.max_row + 1)
               if ws.cell(row, comp_c).value not in (None, "")
               and any(ws.cell(row, headers[n]).value in (None, "")
                       for n in (*group_cols, *fill_in_group_cols))]
    if missing:
        print(f"WARNING: rows still missing values: {missing}")

    wb.save(output_file)
    print(f"Saved: {output_file}")
    return output_file
